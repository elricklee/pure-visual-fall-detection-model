from __future__ import annotations

import argparse
import json
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
DATE_TEXT = "2026年7月28日"
BODY_FONT = "Microsoft YaHei"
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "5F6368"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
GREEN = "1B5E20"
GOLD = "7A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 100, "bottom": 100, "start": 120, "end": 120}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build member-B and project status reports.")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "reports")
    return parser.parse_args()


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _git(*args: str) -> str:
    return subprocess.check_output(
        ["git", *args],
        cwd=ROOT,
        text=True,
        encoding="utf-8",
    ).strip()


def _set_font(run, size: float | None = None, bold: bool | None = None,
              color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_width(parent, tag: str, width: int) -> None:
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    child.set(qn("w:type"), "dxa")
    child.set(qn("w:w"), str(width))


def _apply_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    _set_width(tbl_pr, "w:tblW", CONTENT_WIDTH_DXA)
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for index, width in enumerate(widths):
        table.columns[index].width = Twips(width)
    for row in table.rows:
        row.height = None
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_width(cell._tc.get_or_add_tcPr(), "w:tcW", widths[index])
            _set_cell_margins(cell)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_repeat_keep(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    _set_font(run, size=9, color=GRAY)


def _configure_document(doc: Document, running_label: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run(running_label)
    _set_font(run, size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("华为专项赛道 · 纯视觉跌倒检测MVP  |  第 ")
    _set_font(run, size=9, color=GRAY)
    _add_page_number(footer)
    run = footer.add_run(" 页")
    _set_font(run, size=9, color=GRAY)


def _add_title_block(doc: Document, title: str, subtitle: str, report_id: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    _set_font(run, size=23, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(subtitle)
    _set_font(run, size=13, color=GRAY)

    metadata = [
        ("项目：", "校企协同数字技术大赛 · 华为专项赛道 · 纯视觉跌倒检测MVP"),
        ("责任成员：", "组员乙（端侧部署、性能调优与交叉验收）"),
        ("报告日期：", DATE_TEXT),
        ("报告编号：", report_id),
        ("数据口径：", "本地实测、仓库审计、合成多人功能验证；不包含真实NPU板卡实测"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(label)
        _set_font(label_run, size=10.5, bold=True, color=NAVY)
        value_run = p.add_run(value)
        _set_font(value_run, size=10.5)


def _add_callout(doc: Document, label: str, text: str, fill: str = CALLOUT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    _apply_table_geometry(table, [CONTENT_WIDTH_DXA])
    _repeat_header(table.rows[0])
    cell = table.cell(0, 0)
    _shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(f"{label}：")
    _set_font(run, size=10.5, bold=True, color=NAVY)
    run = p.add_run(text)
    _set_font(run, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    _set_font(run, size=11)


def _add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    _set_font(run, size=11)


def _add_labeled_para(doc: Document, label: str, text: str,
                      color: str = NAVY) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    run = p.add_run(label)
    _set_font(run, bold=True, color=color)
    run = p.add_run(text)
    _set_font(run)


def _add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
    *,
    header_fill: str = LIGHT_GRAY,
    font_size: float = 9.3,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _shade(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        _set_font(run, size=font_size, bold=True, color=NAVY)
    _repeat_header(table.rows[0])

    for row_data in rows:
        row = table.add_row()
        _set_repeat_keep(row)
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            color = None
            if str(value) in {"通过", "已完成", "达标"}:
                color = GREEN
                run.bold = True
            elif str(value) in {"待验证", "待外部条件", "有条件通过"}:
                color = GOLD
                run.bold = True
            elif str(value) in {"未完成", "不通过", "高风险"}:
                color = RED
                run.bold = True
            _set_font(run, size=font_size, color=color)
    _apply_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"数据来源：{text}")
    _set_font(run, size=8.5, color=GRAY, italic=True)


def _add_figure(doc: Document, image_path: Path, caption: str) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(6.15))
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", "多人跌倒检测功能验证截图")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(caption)
    _set_font(run, size=9, color=GRAY, italic=True)


def _page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def _load_context() -> dict:
    fp32_audit = _read_json(ROOT / "reports" / "model" / "fp32_model_audit.json")
    fp32_benchmark = _read_json(
        ROOT / "reports" / "model" / "fp32_cpu_benchmark_reference.json"
    )
    audit = _read_json(ROOT / "reports" / "model" / "int8_model_audit.json")
    benchmark = _read_json(ROOT / "reports" / "model" / "int8_cpu_benchmark.json")
    pose = _read_json(ROOT / "reports" / "evaluation" / "int8_pose_eval_visible.json")
    pseudo_ir = _read_json(ROOT / "reports" / "evaluation" / "int8_pose_eval_pseudo_ir.json")
    visible_video = _read_json(ROOT / "reports" / "evaluation" / "int8_video_eval_visible.json")
    ir_video = _read_json(ROOT / "reports" / "evaluation" / "int8_video_eval_pseudo_ir.json")
    npu = _read_json(ROOT / "reports" / "model" / "int8_npu_feasibility.json")
    multi = _read_json(
        ROOT / "reports" / "evaluation" / "multi_person_functional_validation.json"
    )

    label_paths = list((ROOT / "datasets" / "fall_pose" / "labels").rglob("*.txt"))
    max_persons = 0
    multi_labeled_frames = 0
    for path in label_paths:
        count = sum(1 for line in path.read_text(encoding="utf-8").splitlines() if line.strip())
        max_persons = max(max_persons, count)
        multi_labeled_frames += int(count >= 2)

    model_paths = {
        "fp32": ROOT / "artifacts" / "onnx" / "fall_pose_384_fp32.onnx",
        "int8": ROOT / "artifacts" / "onnx" / "fall_pose_384_int8_qdq.onnx",
        "pt": ROOT / "artifacts" / "pytorch" / "7.26_Train" / "best.pt",
    }
    return {
        "fp32_audit": fp32_audit,
        "fp32_benchmark": fp32_benchmark,
        "audit": audit,
        "benchmark": benchmark,
        "pose": pose,
        "pseudo_ir": pseudo_ir,
        "visible_video": visible_video,
        "ir_video": ir_video,
        "npu": npu,
        "multi": multi,
        "models": {
            key: {
                "path": str(path.relative_to(ROOT)),
                "bytes": path.stat().st_size,
                "mib": round(path.stat().st_size / (1024 * 1024), 3),
            }
            for key, path in model_paths.items()
        },
        "git": {
            "branch": _git("branch", "--show-current"),
            "head": _git("rev-parse", "--short", "HEAD"),
            "subject": _git("log", "-1", "--pretty=%s"),
            "status": _git("status", "--short"),
        },
        "dataset": {
            "labeled_frames": len(label_paths),
            "max_persons_per_frame": max_persons,
            "multi_labeled_frames": multi_labeled_frames,
            "val_images": len(list((ROOT / "datasets" / "fall_pose" / "images" / "val").glob("*.jpg"))),
            "raw_videos": len(list((ROOT / "datasets" / "fall_pose" / "videos" / "raw").glob("*.mp4"))),
        },
        "lap_version": subprocess.check_output(
            [
                str(ROOT / ".venv" / "Scripts" / "python.exe"),
                "-c",
                "from importlib.metadata import version; print(version('lap'))",
            ],
            text=True,
            encoding="utf-8",
        ).strip(),
        "test_count": 133,
    }


def build_member_b_report(context: dict, output_path: Path) -> None:
    doc = Document()
    _configure_document(doc, "组员乙工作完成报告  |  内部验收版")
    _add_title_block(
        doc,
        "组员乙可执行事项完成报告",
        "端侧部署、性能调优、多人功能验收与交付材料",
        "B-WORK-20260728",
    )
    _add_callout(
        doc,
        "完成结论",
        "乙在当前设备和现有仓库条件下能够独立完成的环境补齐、代码审计、多人功能测试、"
        "端到端联调、指标口径修正、告警机制优化和证据归档均已完成。真实同屏多人精度、"
        "具体NPU端延迟及运行内存仍依赖外部素材或目标板卡，不能虚报为实测完成。",
        BLUE_GRAY,
    )

    doc.add_heading("一、职责依据与本次执行范围", level=1)
    _add_labeled_para(
        doc,
        "职责依据：",
        "《三人交叉协作分工方案》将乙定位为“主攻端侧部署与性能调优，交叉参与模型轻量化、"
        "数据测试和文档性能板块撰写”。",
    )
    _add_table(
        doc,
        ["计划阶段", "乙的核心职责", "本次执行结果"],
        [
            ["Day 1", "部署环境、测速工具、指标清单", "补装lap 0.5.13；pip check通过；复核运行依赖"],
            ["Day 2", "基线模型性能评测与问题反馈", "复核模型体积、参数、PC参考速度及多人端到端速度"],
            ["Day 3", "INT8量化、预处理与复杂场景核验", "复用已完成INT8模型；验证可见光、多人成像链路"],
            ["Day 4", "全流程联调与正式性能报告", "完成3轮多人全流程运行、JSON/CSV/HTML/视频留证"],
            ["Day 5", "Demo、部署说明和交付汇总", "补齐多人素材工具、汇总工具、README及两份正式报告"],
        ],
        [1260, 3240, 4860],
    )

    doc.add_heading("二、已完成的具体事项", level=1)
    doc.add_heading("2.1 环境与依赖", level=2)
    _add_bullet(doc, f"在项目虚拟环境安装并验证 ByteTrack 必需依赖 lap {context['lap_version']}。")
    _add_bullet(doc, "执行 pip check，结果为 No broken requirements found。")
    _add_bullet(doc, "验证 INT8 ONNX 可通过 ONNX Runtime CPUExecutionProvider 与 ByteTrack 联合运行。")

    doc.add_heading("2.2 多人评估口径修正", level=2)
    _add_bullet(doc, "将多人模式 person_frames 改为“至少检测到一人的视频帧数”，不再累计人次。")
    _add_bullet(doc, "将 raw_fall_frames、temporal_fall_frames 改为按帧去重；同帧多人跌倒只计1帧。")
    _add_bullet(doc, "新增单元测试，验证两人同帧跌倒时各指标只增加1。")

    doc.add_heading("2.3 告警与运行稳定性优化", level=2)
    _add_bullet(doc, "告警由“持续状态按冷却时间重复触发”改为“进入确认跌倒状态时触发一次”。")
    _add_bullet(doc, "人员恢复后再次跌倒，且满足每ID冷却时间时，才创建新事件。")
    _add_bullet(doc, "Webhook 改为后台发送，并在退出前有限等待，避免网络请求阻塞逐帧推理或丢失请求。")
    _add_bullet(doc, "增加输出视频写入器可用性检查，失败时明确退出。")
    _add_bullet(doc, "逐帧日志增加 detected_person_count 与 untracked_detections，便于定位跟踪丢失。")

    doc.add_heading("2.4 测试工具与证据链", level=2)
    _add_bullet(doc, "新增双视频拼接工具，支持 full/left/right ROI 裁剪，专供多人管线功能测试。")
    _add_bullet(doc, "新增多人运行汇总工具，自动统计同帧人数、ID、未跟踪检测、事件数和端到端FPS。")
    _add_bullet(doc, "生成标注视频、事件截图、回放、frames.jsonl、events.jsonl、run_summary.json 和 HTML 报告。")
    _add_bullet(doc, f"执行 {context['test_count']} 项自动化测试、compileall 和依赖完整性检查，全部通过。")

    _page_break(doc)
    doc.add_heading("三、多人端到端实测结果", level=1)
    runs = context["multi"]["runs"]
    rows = []
    for run in runs:
        model_label = "INT8 ONNX" if str(run["model"]).endswith(".onnx") else "PyTorch FP32"
        rows.append([
            model_label,
            f"{run['image_size']} / {run['confidence_threshold']}",
            f"{run['duration_sec']:.3f}s",
            f"{run['observed_end_to_end_fps']:.3f}",
            str(run["frames_with_2plus_tracked_persons"]),
            ",".join(str(x) for x in run["unique_track_ids"]),
            str(run["event_starts"]),
        ])
    _add_table(
        doc,
        ["模型", "输入/阈值", "56帧耗时", "端到端FPS", "同帧≥2人", "出现ID", "事件"],
        rows,
        [1500, 1280, 1160, 1160, 1180, 1680, 1400],
        font_size=8.7,
    )
    _add_source_note(
        doc,
        "reports/evaluation/multi_person_functional_validation.json；耗时包含视频解码、姿态推理、"
        "ByteTrack、状态机、绘图和文件写入，不能与纯模型推理耗时直接等同。",
    )
    _add_callout(
        doc,
        "实测解释",
        "三轮运行均实际出现同帧2人，且生成独立ID、独立状态、事件截图与回放，证明功能链路通过。"
        "但56帧中稳定同时跟踪2人的帧数仅5至10帧，并出现3至4个不同ID，说明当前模型与ByteTrack"
        "在小目标/姿态剧烈变化时存在ID碎片化，尚不能据此宣称真实多人精度达标。",
    )

    screenshot = (
        ROOT
        / "runs"
        / "system"
        / "multi_person_fp32_roi"
        / "events"
        / "fall_event_p4_frame000054"
        / "snapshot.jpg"
    )
    _add_figure(
        doc,
        screenshot,
        "图1  合成双人功能测试中的独立ID与独立跌倒状态（仅作管线验证）",
    )

    doc.add_heading("四、性能与合规复核", level=1)
    audit = context["audit"]
    bench = context["benchmark"]
    _add_table(
        doc,
        ["项目", "当前结果", "性质", "判断"],
        [
            ["参数量", f"{audit['parameter_count_million']:.4f}M", "ONNX实测", "达标"],
            ["INT8模型文件", f"{context['models']['int8']['mib']:.3f} MiB", "文件实测", "达标"],
            ["PC纯推理均值", f"{bench['latency']['mean_ms']:.4f} ms", "CPU参考", "通过"],
            ["PC纯推理P95", f"{bench['latency']['p95_ms']:.4f} ms", "CPU参考", "通过"],
            ["NPU端延迟≤100ms", "无目标板卡数据", "设计目标", "待验证"],
            ["NPU运行内存≤20MB", "无目标板卡数据", "设计目标", "待验证"],
        ],
        [2200, 2100, 2100, 2960],
    )

    doc.add_heading("五、乙的交付物清单", level=1)
    _add_table(
        doc,
        ["类别", "交付物", "状态"],
        [
            ["代码", "scripts/create_multi_person_fixture.py", "已完成"],
            ["代码", "tools/summarize_multi_validation.py", "已完成"],
            ["整改", "多人评估按帧计数、单次事件触发、异步Webhook、写入器检查", "已完成"],
            ["测试", "tests/test_multi_runtime.py；全套133项测试通过", "已完成"],
            ["原始证据", "reports/evaluation/multi_person_functional_validation.json", "已完成"],
            ["评估结果", "reports/evaluation/multi_person_pipeline_eval.json / .csv", "已完成"],
            ["运行成果", "runs/system/multi_person_*/ 下的视频、事件和HTML", "已完成"],
            ["文档", "乙可执行事项完成报告、项目当前情况详细报告", "已完成"],
        ],
        [1500, 6000, 1860],
    )

    doc.add_heading("六、尚需外部输入的事项", level=1)
    _add_number(doc, "由甲或丙提供至少一段真实同屏2至4人视频，并包含正常、单人跌倒、多人遮挡和交叉行走。")
    _add_number(doc, "为视频提供明确真值：跌倒人员、开始帧、结束帧及正常人员，才能计算准确率、漏报率和误报率。")
    _add_number(doc, "确定目标海思/瑞芯微芯片或提供板卡后，才能执行SDK转换、NPU延迟和运行内存实测。")
    _add_number(doc, "确认告警策略是否允许跟踪ID变化后重复告警；当前按每个track_id的一次跌倒事件处理。")

    _add_callout(
        doc,
        "乙的最终状态",
        "当前可独立工作已完成；多人功能验收为“有条件通过”，真实多人精度与NPU端指标为“待外部条件”。",
        BLUE_GRAY,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def build_project_status_report(context: dict, output_path: Path) -> None:
    doc = Document()
    _configure_document(doc, "项目当前情况详细报告  |  2026-07-28")
    _add_title_block(
        doc,
        "纯视觉跌倒检测MVP项目当前情况详细报告",
        "代码、数据、模型、性能、多人能力、合规性与风险全景",
        "PROJECT-STATUS-20260728",
    )
    _add_callout(
        doc,
        "总体判断",
        "项目已经形成可运行的纯视觉跌倒检测MVP：具备姿态模型、规则评分、时序确认、多人跟踪、"
        "事件告警、截图回放、逐帧日志和HTML报告；INT8模型在参数量与文件体积方面满足轻量化设计。"
        "当前最大缺口不是功能代码，而是真实多人标注数据和具体NPU板卡实测。",
        BLUE_GRAY,
    )

    doc.add_heading("一、项目状态总览", level=1)
    git = context["git"]
    _add_table(
        doc,
        ["维度", "当前状态", "判断"],
        [
            ["代码基线", f"分支 {git['branch']}；HEAD {git['head']}（{git['subject']}）", "稳定基线存在"],
            ["本轮工作区", "多人验收整改、工具和报告已生成，尚未提交Git", "待团队复核提交"],
            ["模型", "FP32 ONNX、INT8 QDQ ONNX、训练后best.pt齐全", "已具备"],
            ["自动化测试", f"{context['test_count']}项全部通过；compileall通过；pip check通过", "通过"],
            ["多人功能", "可同时绘制独立ID/状态并生成独立事件", "有条件通过"],
            ["真实多人精度", "无真实同屏多人标注集", "待验证"],
            ["NPU实测", "未确定具体板卡，未执行厂商SDK转换", "待验证"],
        ],
        [1900, 5480, 1980],
    )

    doc.add_heading("二、系统技术链路", level=1)
    _add_number(doc, "视频、摄像头或流媒体输入；可选灰度复制的伪红外预处理。")
    _add_number(doc, "YOLOv8 Pose输出人体框、检测置信度和17个人体关键点。")
    _add_number(doc, "ByteTrack生成track_id，将同一人员跨帧关联。")
    _add_number(doc, "TrackerManager按track_id分配独立FallDetector和TemporalFallStateMachine。")
    _add_number(doc, "规则评分结合框体宽高比、躯干角度、肩髋距离及关键点质量。")
    _add_number(doc, "状态机完成疑似跌倒、确认跌倒、恢复与误报抑制。")
    _add_number(doc, "系统输出标注视频、事件截图、回放、事件流、逐帧trace、摘要和HTML报告。")
    _add_labeled_para(
        doc,
        "核心优势：",
        "多人状态完全隔离，避免一个人的跌倒投票污染另一个人的状态；事件材料能够追溯到具体ID、"
        "触发帧、分数、原因和回放片段。",
    )

    doc.add_heading("三、代码与工程状态", level=1)
    _add_table(
        doc,
        ["模块", "作用", "当前状态"],
        [
            ["fall_detection/fall_logic.py", "单帧规则评分与内部投票", "已完成"],
            ["fall_detection/state_machine.py", "确认、保持、恢复状态机", "已完成"],
            ["fall_detection/tracker_manager.py", "按track_id隔离每个人的状态", "已完成"],
            ["scripts/infer_video.py", "单/多人标注视频推理", "已完成"],
            ["scripts/evaluate_videos.py", "批量视频评估；多人指标按帧去重", "已整改"],
            ["scripts/run_system.py", "事件中心、告警、截图、回放、trace、报告", "已整改"],
            ["scripts/create_multi_person_fixture.py", "多人功能测试素材生成", "新增完成"],
            ["tools/summarize_multi_validation.py", "多人运行汇总", "新增完成"],
        ],
        [3000, 4380, 1980],
    )

    _page_break(doc)
    doc.add_heading("四、数据集现状", level=1)
    dataset = context["dataset"]
    _add_table(
        doc,
        ["项目", "数量/结果", "解释"],
        [
            ["原始视频", str(dataset["raw_videos"]), "当前本地视频均未通过文件名提供fall/normal真值"],
            ["标注帧", str(dataset["labeled_frames"]), "训练与验证标签总数"],
            ["验证图像", str(dataset["val_images"]), "样本很小，仅适合作回归检查"],
            ["每帧最大人数", str(dataset["max_persons_per_frame"]), "现有标注帧均为单人"],
            ["同帧≥2人标注", str(dataset["multi_labeled_frames"]), "真实多人精度无法据此评估"],
        ],
        [2100, 1800, 5460],
    )
    _add_callout(
        doc,
        "数据风险",
        "现有111个标注帧中同屏多人样本为0。YOLO Pose本身具备多人检测能力，新增ByteTrack与"
        "按ID状态机也能运行，但训练后模型是否能在真实遮挡、交叉走动和多人跌倒场景稳定工作，"
        "目前没有有效数据证据。",
    )

    doc.add_heading("五、模型资产与轻量化状态", level=1)
    audit = context["audit"]
    models = context["models"]
    _add_table(
        doc,
        ["模型资产", "文件体积", "用途", "状态"],
        [
            ["fall_pose_384_fp32.onnx", f"{models['fp32']['mib']:.3f} MiB", "通用FP32 ONNX", "已完成"],
            ["fall_pose_384_int8_qdq.onnx", f"{models['int8']['mib']:.3f} MiB", "PC兼容QDQ INT8", "已完成"],
            ["7.26_Train/best.pt", f"{models['pt']['mib']:.3f} MiB", "训练后PyTorch模型", "已完成"],
        ],
        [4200, 1500, 2100, 1560],
    )
    _add_labeled_para(doc, "参数规模：", f"{audit['parameter_count_million']:.4f}M，低于赛题20M上限。")
    _add_labeled_para(doc, "INT8文件：", f"{audit['file_mb']:.4f} MB（十进制）/ {models['int8']['mib']:.3f} MiB。")
    _add_labeled_para(
        doc,
        "算子审计：",
        "ONNX opset 13；主要算子包括Conv、Concat、Resize、Sigmoid、Softmax、"
        "QuantizeLinear和DequantizeLinear。量化算子仍需目标SDK确认支持。",
    )

    doc.add_heading("六、已有精度与速度数据", level=1)
    pose = context["pose"]["metrics"]
    ir_pose = context["pseudo_ir"]["metrics"]
    bench = context["benchmark"]
    _add_table(
        doc,
        ["数据/指标", "当前结果", "性质与限制"],
        [
            ["可见光框 mAP50", f"{pose['metrics/mAP50(B)']:.3f}", "11张验证图像，小样本回归"],
            ["可见光姿态 mAP50", f"{pose['metrics/mAP50(P)']:.3f}", "11张验证图像，小样本回归"],
            ["可见光姿态 mAP50-95", f"{pose['metrics/mAP50-95(P)']:.4f}", "不能代表正式比赛泛化精度"],
            ["伪红外姿态 mAP50", f"{ir_pose['metrics/mAP50(P)']:.3f}", "灰度模拟，不是真实红外相机"],
            ["伪红外姿态 mAP50-95", f"{ir_pose['metrics/mAP50-95(P)']:.4f}", "存在较明显下降"],
            ["PC纯推理均值", f"{bench['latency']['mean_ms']:.4f} ms", "100次；不含解码/NMS/跟踪/绘图"],
            ["PC纯推理P95", f"{bench['latency']['p95_ms']:.4f} ms", "CPU参考，不是NPU实测"],
        ],
        [2700, 1900, 4760],
    )
    _add_labeled_para(
        doc,
        "视频级标签问题：",
        "10段可见光原视频与10段伪红外处理视频的expected_fall均为null，因此历史报告只能记录"
        "是否触发告警，TP/FP/TN/FN和accuracy均无法计算。",
        RED,
    )

    doc.add_heading("七、多人功能现状", level=1)
    multi_rows = []
    for run in context["multi"]["runs"]:
        label = "INT8 ONNX" if str(run["model"]).endswith(".onnx") else "PyTorch FP32"
        multi_rows.append([
            label,
            f"{run['image_size']}",
            f"{run['observed_end_to_end_fps']:.3f}",
            f"{run['frames_with_2plus_tracked_persons']}/56",
            ",".join(map(str, run["unique_track_ids"])),
            str(run["event_starts"]),
        ])
    _add_table(
        doc,
        ["模型", "imgsz", "端到端FPS", "同帧≥2人", "出现ID", "事件"],
        multi_rows,
        [1900, 1100, 1600, 1700, 2000, 1060],
    )
    _add_callout(
        doc,
        "多人结论",
        "功能层面已通过：可以同帧检测两人、分别编号、独立判定并分别留证。性能层面仍有条件："
        "合成视频中第二人的持续检出不足，track_id存在重建；降低INT8置信阈值至0.15未形成稳定改善。"
        "真实多人数据到位前，不应在PPT中写“多人准确率已验证”。",
    )

    _page_break(doc)
    doc.add_heading("八、赛题硬指标与证据边界", level=1)
    _add_table(
        doc,
        ["赛题/内部指标", "当前证据", "结论"],
        [
            ["参数量≤20M", f"{audit['parameter_count_million']:.4f}M，ONNX实测", "达标"],
            [
                "FP32权重≤80MB",
                f"{context['fp32_audit']['initializer_mb']:.4f}MB，FP32 ONNX实测",
                "达标",
            ],
            ["INT8权重目标≤20MB", f"INT8文件{models['int8']['mib']:.3f}MiB", "达标"],
            ["端侧推理≤100ms", f"PC纯推理均值{bench['latency']['mean_ms']:.4f}ms；无目标NPU数据", "待验证"],
            ["NPU运行内存≤20MB", "只有权重存储估算；无运行时峰值", "待验证"],
            ["目标SDK算子兼容", "Q/DQ算子需海思或瑞芯微转换日志确认", "待验证"],
        ],
        [2600, 4700, 2060],
    )
    _add_source_note(
        doc,
        "reports/model/int8_model_audit.json、reports/model/int8_cpu_benchmark.json、"
        "reports/model/int8_npu_feasibility.json。",
    )

    doc.add_heading("九、质量保证与可复现性", level=1)
    _add_bullet(doc, f"自动化测试：{context['test_count']}项全部通过。")
    _add_bullet(doc, "Python compileall通过；pip check显示无破损依赖。")
    _add_bullet(doc, f"ByteTrack运行依赖：lap {context['lap_version']} 已安装。")
    _add_bullet(doc, "三次多人运行均保留完整参数、逐帧trace、事件JSONL、截图、回放和HTML报告。")
    _add_bullet(doc, "多人评估输出同时提供JSON和CSV，便于PPT、统计表和后续自动对比。")
    _add_bullet(doc, "工作区尚未提交，便于团队先复核差异后再形成正式Git版本。")

    doc.add_heading("十、主要风险与优先级", level=1)
    _add_table(
        doc,
        ["优先级", "风险", "影响", "建议动作"],
        [
            ["P0", "缺少真实同屏多人标注数据", "无法证明多人精度、漏报和ID稳定性", "采集并逐人标注2至4人视频"],
            ["P0", "无具体NPU板卡与SDK转换记录", "100ms与20MB不能实测闭环", "确定芯片后执行转换和板端测试"],
            ["P1", "多人场景ID碎片化", "同一事件可能因ID变化被拆分", "调ByteTrack阈值并补遮挡/交叉样本"],
            ["P1", "INT8端到端链路较慢", "纯推理快不等于完整Demo快", "剥离Ultralytics后处理并分阶段测速"],
            ["P1", "视频缺少fall/normal真值", "历史视频报告accuracy为空", "建立视频级与事件级标注清单"],
            ["P2", "Webhook未连接真实业务端", "仅验证代码路径，未验证接收方", "接入测试端点并做超时/重试测试"],
        ],
        [900, 2600, 2860, 3000],
        font_size=8.8,
    )

    doc.add_heading("十一、建议的下一轮执行顺序", level=1)
    _add_number(doc, "采集或获取真实多人视频，建立每人、每事件和时间段真值。")
    _add_number(doc, "在PyTorch与INT8 ONNX上分别跑同一批素材，定位检出、跟踪、规则或量化造成的差异。")
    _add_number(doc, "调优ByteTrack参数并验证遮挡、交叉、进出画面和ID重建。")
    _add_number(doc, "把端到端耗时拆分为解码、预处理、模型、NMS、跟踪、状态机、绘图和写盘。")
    _add_number(doc, "确定目标NPU后执行SDK转换、算子回退检查、真实延迟和内存测试。")
    _add_number(doc, "将最终实测数据、限制说明和演示视频同步到PPT与答辩材料。")

    doc.add_heading("十二、可用于提交材料的严谨表述", level=1)
    _add_callout(
        doc,
        "推荐表述",
        "当前系统已完成多人跟踪与按ID独立跌倒判定的功能验证，并可自动输出告警截图、回放和结构化日志。"
        "INT8模型参数量与权重体积满足轻量化设计目标；PC端纯推理速度为参考值。真实多人精度、"
        "目标NPU端延迟与运行内存将在获得标准多人测试集和具体板卡后进一步实测。",
        BLUE_GRAY,
    )
    _add_labeled_para(
        doc,
        "禁止表述：",
        "“多人准确率已达标”“海思/瑞芯微NPU实测≤100ms”“NPU内存实测≤20MB”。"
        "当前均无相应真实数据或设备证据。",
        RED,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    args = parse_args()
    context = _load_context()
    member_b_path = args.output_dir / "组员乙可执行事项完成报告_20260728.docx"
    project_path = args.output_dir / "纯视觉跌倒检测MVP项目当前情况详细报告_20260728.docx"
    build_member_b_report(context, member_b_path)
    build_project_status_report(context, project_path)
    print(member_b_path)
    print(project_path)


if __name__ == "__main__":
    main()

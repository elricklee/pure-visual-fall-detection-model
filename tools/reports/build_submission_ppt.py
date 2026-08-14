from __future__ import annotations

import json
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

try:
    from PIL import Image
except Exception:  # pragma: no cover - python-pptx normally installs Pillow
    Image = None


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "deliverables" / "ppt"
PPTX = OUT_DIR / "fall_detection_mvp_draft.pptx"
CHECKLIST = OUT_DIR / "ppt_assets_checklist.md"

RED = RGBColor(199, 0, 11)
DARK = RGBColor(26, 26, 26)
MID = RGBColor(84, 84, 84)
LIGHT = RGBColor(245, 246, 248)
LINE = RGBColor(222, 226, 232)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(18, 134, 89)
BLUE = RGBColor(33, 102, 172)
ORANGE = RGBColor(214, 106, 28)


def load_json(relative: str) -> dict:
    path = ROOT / relative
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8-sig"))


eval_report = load_json("reports/evaluation/pseudo_ir_labeled_video_eval.json")
quant_report = load_json("reports/model/int8_quantization.json")
bench_report = load_json("reports/model/int8_cpu_benchmark.json")
compliance_report = load_json("reports/compliance/model_compliance_report.json")

summary = eval_report.get("summary", {})
tp = int(summary.get("TP", 0))
fp = int(summary.get("FP", 0))
tn = int(summary.get("TN", 0))
fn = int(summary.get("FN", 0))
total = tp + fp + tn + fn
accuracy = float(summary.get("accuracy", 0.0))
precision = tp / (tp + fp) if (tp + fp) else 0
recall = tp / (tp + fn) if (tp + fn) else 0
compression = float(quant_report.get("compression_ratio", 0.0))
latency = bench_report.get("latency", {})
mean_ms = float(latency.get("mean_ms", 0.0))
fps = float(latency.get("fps_from_mean", 0.0))
comp_sum = compliance_report.get("summary", {})


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(path).replace("\\", "/")


def add_box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=18,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    font="Microsoft YaHei",
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_multiline(slide, x, y, w, h, lines, size=16, color=DARK, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for idx, line in enumerate(lines):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
        if bullet:
            p.text = "• " + line
    return box


def add_title(slide, page_no, title, subtitle=None):
    add_text(slide, 0.45, 0.23, 0.7, 0.3, f"{page_no:02d}", 12, RED, True)
    add_text(slide, 1.05, 0.19, 8.9, 0.45, title, 25, DARK, True)
    if subtitle:
        add_text(slide, 1.07, 0.66, 9.8, 0.28, subtitle, 10, MID)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.93), Inches(12.45), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = RED
    line.line.fill.background()


def add_footer(slide, note):
    add_text(slide, 0.45, 7.14, 10.8, 0.22, "证据口径：" + note, 8.5, MID)
    add_text(slide, 12.05, 7.14, 0.85, 0.22, "Draft", 8.5, RED, True, PP_ALIGN.RIGHT)


def add_metric(slide, x, y, w, h, value, label, color=RED, note=None):
    add_box(slide, x, y, w, h)
    add_text(slide, x + 0.16, y + 0.14, w - 0.32, 0.42, value, 27, color, True)
    add_text(slide, x + 0.16, y + 0.68, w - 0.32, 0.25, label, 10.5, DARK, True)
    if note:
        add_text(slide, x + 0.16, y + 0.98, w - 0.32, 0.24, note, 8.5, MID)


def add_placeholder(slide, x, y, w, h, title, detail):
    box = add_box(slide, x, y, w, h, RGBColor(252, 248, 248), RED)
    add_text(slide, x + 0.22, y + 0.24, w - 0.44, 0.35, title, 14, RED, True)
    add_multiline(slide, x + 0.22, y + 0.72, w - 0.44, h - 0.9, detail, 11, MID)
    return box


def add_image(slide, relative_path, x, y, w, h, label):
    path = ROOT / relative_path
    add_box(slide, x, y, w, h, RGBColor(250, 250, 250), LINE)
    if not path.exists() or Image is None:
        add_placeholder(slide, x, y, w, h, "待补图片", [label, relative_path])
        return False
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    add_text(slide, x + 0.12, y + h - 0.28, w - 0.24, 0.2, label, 8, MID)
    return True


def add_flow(slide, labels, x, y, w, h, color=RED):
    gap = 0.18
    node_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        nx = x + i * (node_w + gap)
        add_box(slide, nx, y, node_w, h, RGBColor(255, 255, 255), color)
        add_text(slide, nx + 0.08, y + 0.18, node_w - 0.16, h - 0.3, label, 11, DARK, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < len(labels) - 1:
            add_text(slide, nx + node_w + 0.02, y + 0.28, 0.14, 0.25, ">", 13, color, True, PP_ALIGN.CENTER)


def add_bars(slide, x, y, w, h, rows, color=RED):
    max_v = max(v for _, v, _ in rows) if rows else 1
    for i, (label, value, value_label) in enumerate(rows):
        yy = y + i * (h / len(rows))
        add_text(slide, x, yy, 1.8, 0.28, label, 10.5, DARK, True)
        add_box(slide, x + 1.85, yy + 0.05, w - 2.65, 0.18, RGBColor(238, 240, 244), RGBColor(238, 240, 244), False)
        bar_w = (w - 2.65) * (value / max_v)
        add_box(slide, x + 1.85, yy + 0.05, bar_w, 0.18, color, color, False)
        add_text(slide, x + w - 0.7, yy - 0.02, 0.7, 0.25, value_label, 9.5, MID, True, PP_ALIGN.RIGHT)


def init_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    return slide


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1
    s = init_slide(prs)
    add_box(s, 0, 0, 13.333, 7.5, WHITE, WHITE, False)
    add_box(s, 0, 0, 0.22, 7.5, RED, RED, False)
    add_text(s, 0.65, 0.75, 10.7, 0.35, "纯视觉跌倒检测 MVP", 16, RED, True)
    add_text(s, 0.65, 1.28, 8.8, 1.25, "不只是单点模型，而是端侧告警闭环", 37, DARK, True)
    add_multiline(
        s,
        0.7,
        2.75,
        5.7,
        1.0,
        ["YOLO Pose + 几何评分 + 时序状态机", "ONNX/INT8 端侧部署准备", "输出视频、截图、JSON 告警证据"],
        15,
        MID,
    )
    add_image(s, "samples/reference_output_fp32.jpg", 7.1, 1.05, 5.4, 3.8, "建议替换为现场 Demo 首帧/系统运行截图")
    add_placeholder(s, 7.1, 5.1, 5.4, 1.0, "封面图待增强", ["可放团队现场演示照片", "或摄像头 + 边缘设备 + 检测画面的合成截图"])
    add_footer(s, "封面图片来自 samples/reference_output_fp32.jpg，最终建议替换成现场演示图。")

    # 2
    s = init_slide(prs)
    add_title(s, 2, "痛点明确：低成本、非接触、端侧实时告警", "评委先看应用价值，再看模型指标")
    add_metric(s, 0.7, 1.35, 2.3, 1.28, "纯视觉", "无需可穿戴", RED)
    add_metric(s, 3.35, 1.35, 2.3, 1.28, "端侧", "隐私与弱网可用", BLUE)
    add_metric(s, 6.0, 1.35, 2.3, 1.28, "实时", "事件级告警", GREEN)
    add_metric(s, 8.65, 1.35, 2.3, 1.28, "证据", "截图/视频/JSON", ORANGE)
    add_multiline(
        s,
        0.85,
        3.2,
        5.1,
        2.0,
        ["老人居家和养老场景需要低打扰监测", "摄像头方案部署门槛低，适合边缘设备落地", "端侧处理降低视频外传压力，告警更快"],
        16,
        DARK,
        True,
    )
    add_placeholder(s, 6.6, 3.0, 5.7, 2.5, "场景图片待补", ["建议放：居家/养老院监控场景图", "证据：比赛赛题需求截图或应用场景页"])
    add_footer(s, "本页是场景价值页，需要补赛题原文或需求截图作为外部证据。")

    # 3
    s = init_slide(prs)
    add_title(s, 3, "技术主线：YOLO Pose + 几何评分 + 时序状态机", "主干成立：先看见人，再理解姿态，再做事件级确认")
    add_flow(s, ["视频帧", "YOLO Pose", "关键点", "几何评分", "时序状态机", "告警证据"], 0.75, 1.45, 11.8, 0.9)
    add_multiline(
        s,
        0.85,
        2.9,
        5.2,
        2.3,
        ["检测层：复用成熟姿态估计主干，避免从零训练视觉大模型", "规则层：用身体长宽比、躯干角度、头部高度等姿态特征解释跌倒", "时序层：连续风险而非单帧触发，降低偶发姿态误判"],
        14,
        DARK,
        True,
    )
    add_image(s, "samples/reference_output_fp32.jpg", 6.55, 2.45, 5.5, 3.3, "关键点检测样例")
    add_footer(s, "证据：fall_detection/fall_logic.py、fall_detection/state_machine.py、samples/reference_output_fp32.jpg。")

    # 4
    s = init_slide(prs)
    add_title(s, 4, "可解释算法：每次告警都有姿态依据", "比赛答辩不只讲 AUC，还要讲为什么判跌倒")
    add_multiline(
        s,
        0.85,
        1.38,
        4.7,
        2.5,
        ["躯干角度接近水平：身体姿态异常", "人体框宽高比增大：由站立转为横向", "头部/肩部高度下降：与跌倒过程一致", "置信度和关键点缺失会影响最终风险"],
        15,
        DARK,
        True,
    )
    add_image(s, "runs/debug_frames_state/fall02_state_frame_064.jpg", 6.1, 1.25, 5.7, 3.9, "FALL_CONFIRMED 调试帧")
    add_placeholder(s, 0.85, 4.45, 4.7, 1.05, "建议补图", ["放一张带 score 分解的调试截图", "把角度、宽高比、关键点置信度标在画面旁"])
    add_footer(s, "证据：fall_detection/fall_logic.py 与 debug_frames_state/fall02_state_frame_064.jpg。")

    # 5
    s = init_slide(prs)
    add_title(s, 5, "时序状态机降低单帧误报", "把瞬时风险转成候选、确认、恢复的事件过程")
    add_flow(s, ["NORMAL", "FALL_CANDIDATE", "FALL_CONFIRMED", "RECOVERING", "NORMAL"], 0.75, 1.35, 11.3, 0.9, BLUE)
    add_image(s, "runs/debug_frames_state/fall02_state_frame_060.jpg", 0.85, 2.6, 5.2, 3.3, "状态迁移前后样例")
    add_image(s, "runs/debug_frames_state/fall02_state_frame_064.jpg", 6.5, 2.6, 5.2, 3.3, "确认告警样例")
    add_footer(s, "证据：fall_detection/state_machine.py；调试帧展示从风险积累到确认告警。")

    # 6
    s = init_slide(prs)
    add_title(s, 6, "多人能力：按 track_id 独立维护状态", "架构上支持多人同屏，但真实多人同屏截图还需要补")
    add_box(s, 0.85, 1.35, 3.0, 3.4)
    add_text(s, 1.05, 1.62, 2.6, 0.4, "Person #12", 17, RED, True, PP_ALIGN.CENTER)
    add_multiline(s, 1.05, 2.3, 2.55, 1.4, ["状态：NORMAL", "风险：低", "独立缓存：有"], 13, DARK, True)
    add_box(s, 4.25, 1.35, 3.0, 3.4)
    add_text(s, 4.45, 1.62, 2.6, 0.4, "Person #31", 17, BLUE, True, PP_ALIGN.CENTER)
    add_multiline(s, 4.45, 2.3, 2.55, 1.4, ["状态：CANDIDATE", "风险：上升", "独立缓存：有"], 13, DARK, True)
    add_placeholder(s, 7.65, 1.35, 4.3, 3.4, "真实多人截图待补", ["当前只有功能验证报告", "最终请补一张多人同屏 Demo 截图", "不要把功能验证说成真实场景准确率"])
    add_metric(s, 0.85, 5.28, 2.9, 0.82, "track_id", "多人状态隔离键", RED)
    add_metric(s, 4.25, 5.28, 2.9, 0.82, "functional", "功能验证口径", BLUE)
    add_footer(s, "证据：fall_detection/tracker_manager.py、reports/evaluation/multi_person_functional_validation.json。")

    # 7
    s = init_slide(prs)
    add_title(s, 7, "系统闭环：不只检测，还能生成告警证据", "从模型输出走到可复盘事件，这是工程完整性的核心优势")
    add_flow(s, ["run_system.py", "检测/跟踪", "状态机", "事件目录", "snapshot/video/json"], 0.8, 1.35, 11.3, 0.85, GREEN)
    add_image(s, "runs/system/demo_urfall_fall_01_v2/events/demo_urfall_fall_01_v2_p13_000133/snapshot.jpg", 0.85, 2.55, 5.3, 3.3, "系统事件 snapshot")
    add_multiline(
        s,
        6.65,
        2.7,
        4.8,
        2.25,
        ["告警不是一个布尔值，而是一组事件材料", "输出目录保留截图、视频片段和元数据", "适合答辩现场展示：可定位、可复盘、可审计"],
        15,
        DARK,
        True,
    )
    add_footer(s, "证据：scripts/runtime/run_system.py 与 runs/system/.../snapshot.jpg。")

    # 8
    s = init_slide(prs)
    add_title(s, 8, "评估结果：当前样本上跌倒识别链路成立", "小样本验证通过，但还不是最终泛化结论")
    add_metric(s, 0.8, 1.35, 2.25, 1.15, f"{accuracy * 100:.1f}%", "视频级 Accuracy", RED)
    add_metric(s, 3.35, 1.35, 2.25, 1.15, f"{precision * 100:.0f}%", "Precision", BLUE)
    add_metric(s, 5.9, 1.35, 2.25, 1.15, f"{recall * 100:.0f}%", "Recall", GREEN)
    add_metric(s, 8.45, 1.35, 2.25, 1.15, f"{total}", "标注视频数", ORANGE)
    add_bars(
        s,
        0.95,
        3.05,
        5.6,
        1.65,
        [("TP", tp, str(tp)), ("FP", fp, str(fp)), ("TN", tn, str(tn)), ("FN", fn, str(fn))],
        RED,
    )
    add_multiline(
        s,
        7.0,
        3.0,
        4.9,
        1.5,
        ["当前结果说明完整链路可跑通", "主要短板是非跌倒视频 FP=1", "AUC/阈值优化仍是下一步冲刺点"],
        14,
        DARK,
        True,
    )
    add_placeholder(s, 7.0, 4.78, 4.9, 1.0, "证据截图待补", ["建议放 evaluation report.html / 命令行评估截图", "展示 confusion matrix 和 AUC 曲线"])
    add_footer(s, "证据：reports/evaluation/pseudo_ir_labeled_video_eval.json；8 条视频为小样本。")

    # 9
    s = init_slide(prs)
    add_title(s, 9, "视频级评分优化：从单帧峰值转向时序风险", "当前 AUC 未作为达标结论；建议用时间窗口特征继续拉开正负样本")
    add_multiline(
        s,
        0.85,
        1.35,
        5.5,
        2.35,
        ["已完成：视频级指标脚本和伪红外标注视频评估", "问题：单帧 max_score 容易把部分 ADL 拉高", "改进：使用高风险持续时长、首次告警帧、状态机确认比例做融合分数", "目标：让正样本保持高分，同时压低短暂异常动作"],
        14,
        DARK,
        True,
    )
    add_bars(
        s,
        6.75,
        1.55,
        5.1,
        1.95,
        [("max_score", 0.72, "已用"), ("temporal frames", 0.88, "建议强化"), ("confirmed ratio", 0.80, "建议加入"), ("first alert", 0.55, "辅助")],
        BLUE,
    )
    add_placeholder(s, 6.75, 4.15, 5.1, 1.2, "AUC 图待补", ["放 ROC 曲线 / PR 曲线", "证据文件建议：reports/evaluation/video_auc_report.json 或 report.html"])
    add_footer(s, "口径：本页是改进路线，不宣称当前 AUC 已达到最终目标。")

    # 10
    s = init_slide(prs)
    add_title(s, 10, "端侧部署准备：ONNX + INT8 量化 + 固定输入", "模型侧已经具备迁移到端侧推理框架的工程基础")
    add_metric(s, 0.85, 1.35, 2.65, 1.1, "384x384", "固定输入", RED)
    add_metric(s, 3.85, 1.35, 2.65, 1.1, "QDQ", "INT8 格式", BLUE)
    add_metric(s, 6.85, 1.35, 2.65, 1.1, f"{compression * 100:.2f}%", "INT8/FP32 体积", GREEN)
    add_metric(s, 9.85, 1.35, 2.05, 1.1, "Conv", "量化算子", ORANGE)
    add_flow(s, ["FP32 ONNX", "校准样本", "静态量化", "INT8 QDQ", "模型审计"], 0.9, 3.25, 11.3, 0.82)
    add_multiline(
        s,
        1.0,
        4.65,
        10.7,
        0.7,
        ["伪红外样本是灰度复制流程验证，不等价于真实红外传感器数据；答辩时应主动说明。"],
        13,
        MID,
    )
    add_footer(s, "证据：reports/model/int8_quantization.json、fp32_model_audit.json、int8_model_audit.json。")

    # 11
    s = init_slide(prs)
    add_title(s, 11, "性能与合规证据：已有 CPU 参考，NPU 待实测", "把已经测到的和仍待验证的边界讲清楚")
    add_metric(s, 0.85, 1.35, 2.65, 1.1, f"{mean_ms:.4f}ms", "CPU ORT 平均推理", RED)
    add_metric(s, 3.85, 1.35, 2.65, 1.1, f"{fps:.1f}", "FPS from mean", BLUE)
    add_metric(s, 6.85, 1.35, 2.65, 1.1, str(comp_sum.get("pass", 0)), "合规通过项", GREEN)
    add_metric(s, 9.85, 1.35, 2.05, 1.1, str(comp_sum.get("design_target", 0)), "设计目标项", ORANGE)
    add_multiline(
        s,
        0.95,
        3.1,
        5.4,
        2.2,
        ["CPU 参考测速不包含视频解码、缩放、NMS、时序判别和告警输出", "NPU 内存和端到端延迟当前是 design_target", "正式冲刺应补真实板卡日志、端到端耗时截图和 report.html"],
        14,
        DARK,
        True,
    )
    add_placeholder(s, 6.95, 3.0, 4.85, 2.15, "NPU 证据待补", ["放目标板卡运行日志截图", "放端到端 latency 表格", "放模型转换工具通过截图"])
    add_footer(s, "证据：reports/model/int8_cpu_benchmark.json、reports/compliance/model_compliance_report.json。")

    # 12
    s = init_slide(prs)
    add_title(s, 12, "总结：工程闭环，而不是单一指标", "第一名叙事建议：价值明确、链路完整、证据可复盘、边界诚实")
    add_multiline(
        s,
        0.9,
        1.35,
        5.35,
        2.8,
        ["优势 1：主干简单可靠，姿态估计 + 可解释规则适合端侧", "优势 2：从视频输入到告警事件输出形成闭环", "优势 3：已有量化、审计、合规和评估报告支撑", "短板：样本少、真实红外/NPU/多人场景还需补证据"],
        15,
        DARK,
        True,
    )
    add_box(s, 6.9, 1.35, 4.7, 2.8, RGBColor(255, 255, 255), RED)
    add_text(s, 7.15, 1.7, 4.2, 0.5, "答辩收束句", 19, RED, True, PP_ALIGN.CENTER)
    add_text(
        s,
        7.25,
        2.45,
        4.0,
        0.9,
        "我们交付的是一个可解释、可量化、可部署准备的纯视觉跌倒告警 MVP。",
        20,
        DARK,
        True,
        PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_placeholder(s, 1.0, 5.0, 10.6, 0.9, "最终冲刺证据", ["补 3 张图：现场 Demo、ROC/AUC、NPU/端到端日志；PPT 就能从“能跑”升级到“可信”。"])
    add_footer(s, "本页用于统一团队口径：不夸大未实测内容，把优势落在工程闭环和证据链。")

    return prs


def build_checklist():
    rows = [
        ("1", "项目定位封面", "现场 Demo / 摄像头 + 检测画面", "samples/reference_output_fp32.jpg", "建议替换为团队现场图"),
        ("2", "应用痛点", "居家/养老院场景、赛题需求截图", "README.md / 赛题材料", "待补外部场景证据"),
        ("3", "技术主线", "算法流程图 + 姿态检测样例", "fall_detection/fall_logic.py; samples/reference_output_fp32.jpg", "已放样例图"),
        ("4", "可解释算法", "FALL_CONFIRMED 调试帧 + score 分解", "runs/debug_frames_state/fall02_state_frame_064.jpg", "建议补 score 可视化"),
        ("5", "时序状态机", "NORMAL 到 FALL_CONFIRMED 对比帧", "fall_detection/state_machine.py; runs/debug_frames_state/", "已放两帧"),
        ("6", "多人能力", "真实多人同屏检测截图", "fall_detection/tracker_manager.py; reports/evaluation/multi_person_functional_validation.json", "必须补真实截图，当前只写功能验证"),
        ("7", "系统闭环", "告警事件 snapshot、event.json、视频片段目录截图", "scripts/runtime/run_system.py; runs/system/.../snapshot.jpg", "建议补目录结构截图"),
        ("8", "评估结果", "评估报告截图/混淆矩阵", "reports/evaluation/pseudo_ir_labeled_video_eval.json", "建议补 report.html 或命令行截图"),
        ("9", "AUC 优化", "ROC/PR 曲线、阈值扫描表", "reports/evaluation/*auc* 或后续生成文件", "当前留占位，不宣称达标"),
        ("10", "端侧部署", "ONNX/INT8 模型审计图、模型体积对比", "reports/model/int8_quantization.json", "可补 Netron 截图"),
        ("11", "性能合规", "CPU benchmark 截图、NPU 日志截图", "reports/model/int8_cpu_benchmark.json; reports/compliance/model_compliance_report.json", "NPU 实测待补"),
        ("12", "总结", "三证据拼图：Demo/AUC/NPU", "docs/SUBMISSION_READINESS_REVIEW.md", "终版前补齐"),
    ]
    lines = [
        "# PPT 初版图片与证据清单",
        "",
        "> 口径提醒：伪红外是灰度复制流程验证，不是真实红外传感器；CPU 19.0201ms 是 ONNX Runtime 参考，不是 NPU；8 个标注视频是小样本。",
        "",
        "| 页码 | 页面 | 建议图片/视觉 | 当前证据路径 | TODO |",
        "|---|---|---|---|---|",
    ]
    lines.extend(f"| {p} | {title} | {visual} | `{evidence}` | {todo} |" for p, title, visual, evidence, todo in rows)
    lines.extend(
        [
            "",
            "## 终版最关键的 4 张证据图",
            "",
            "1. 现场端到端 Demo 截图：摄像头画面、检测框、状态、告警时间。",
            "2. ROC/AUC 或阈值扫描图：证明 AUC/阈值优化不是口头说法。",
            "3. 真实多人同屏截图：证明 track_id 独立状态不是只停留在代码功能验证。",
            "4. 目标板卡或 NPU 工具日志：把 design_target 升级成 measured。",
            "",
        ]
    )
    CHECKLIST.write_text("\n".join(lines), encoding="utf-8")


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(PPTX)
    build_checklist()

    with zipfile.ZipFile(PPTX) as zf:
        slide_count = len([n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")])
    print(f"PPTX: {rel(PPTX)}")
    print(f"Checklist: {rel(CHECKLIST)}")
    print(f"Slides: {slide_count}")
    print(f"Bytes: {PPTX.stat().st_size}")


if __name__ == "__main__":
    main()
